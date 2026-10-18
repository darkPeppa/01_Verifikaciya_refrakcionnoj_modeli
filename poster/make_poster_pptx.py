#!/usr/bin/env python3
"""Editable A1 poster (PowerPoint, one slide 594 x 841 mm) from the computed results.

Texts, tables and charts are native PowerPoint objects (charts keep their data: right click ->
"Edit data"); the ray pictures are images exported by make_poster.py --assets.

usage: make_poster_pptx.py --results <stage-1 results> --data <stage-2 data> --assets <dir with PNGs> --out poster_A1.pptx
requires: python-pptx (pip install python-pptx)
"""
import argparse
import csv
import gzip
import json
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Mm, Pt

W, H = Mm(594), Mm(841)
FONT = "Arial"
NAVY, ACC, ORANGE, GREEN, GREY, VIOLET, SBRED = "14325C", "C0392B", "E67E22", "2E8B57", "8C8C8C", "8E44AD", "7A1F1F"
LIGHT, WHITE, TEXT, MUTED = "EEF3F9", "FFFFFF", "1A1A1A", "555555"
NAME = {"objective1": "Объектив 1", "objective2": "Объектив 2", "objective2_N145": "Объектив 2, N=145",
        "objective2_N180": "Объектив 2, N=180", "condenser1": "Конденсор 1", "condenser2": "Конденсор 2"}


def rgb(h):
    return RGBColor.from_string(h)


def rc(x, nd):
    return f"{x:.{nd}f}".replace(".", ",")


def rcsv(path):
    op = gzip.open if path.endswith(".gz") else open
    with op(path, "rt") as f:
        rows = list(csv.DictReader(line for line in f if not line.startswith("#")))
    out = []
    for r in rows:
        o = {}
        for k, v in r.items():
            try:
                o[k] = float(v)
            except (TypeError, ValueError):
                o[k] = v
        out.append(o)
    return out


# ------------------------------------------------------------------ geometry helpers (figure fractions, y from the bottom)
def box_emu(x, y, w, h):
    return Emu(int(x * W)), Emu(int((1 - y - h) * H)), Emu(int(w * W)), Emu(int(h * H))


def rect(slide, x, y, w, h, fill, line=None, rounded=True, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, *box_emu(x, y, w, h))
    if rounded:
        shp.adjustments[0] = 0.02 if w < 0.5 else 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(2.5)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def textbox(slide, x, y_top, w, h, paras, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            name=None, spacing=1.15):
    """paras: list of strings or list of lists of (text, {opts}) runs."""
    tb = slide.shapes.add_textbox(*box_emu(x, y_top - h, w, h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Mm(1)
    tf.margin_top = tf.margin_bottom = Mm(0.5)
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        runs = p if isinstance(p, list) else [(p, {})]
        for t, o in runs:
            r = para.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = rgb(o.get("color", color))
    return tb


def section(slide, x, y, w, h, title, color=NAVY):
    rect(slide, x, y, w, h, WHITE, line=color, name=f"Панель: {title}")
    hdr = rect(slide, x, y + h - 0.021, w, 0.021, color, name=f"Заголовок: {title}")
    tf = hdr.text_frame
    tf.margin_left = Mm(8)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title
    r.font.name = FONT
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = rgb(WHITE)


def table(slide, x, y, w, h, rows, size=14, col_w=None, name=None):
    gs = slide.shapes.add_table(len(rows), len(rows[0]), *box_emu(x, y, w, h))
    if name:
        gs.name = name
    t = gs.table
    total = sum(col_w) if col_w else len(rows[0])
    for j in range(len(rows[0])):
        t.columns[j].width = Emu(int(w * W * (col_w[j] if col_w else 1) / total))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = str(val)
            c.margin_left = c.margin_right = Mm(1)
            c.margin_top = c.margin_bottom = Mm(0.6)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(NAVY if i == 0 else ("F3F6FA" if i % 2 == 0 else WHITE))
            for p in c.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.name = FONT
                    r.font.size = Pt(size)
                    r.font.bold = i == 0
                    r.font.color.rgb = rgb(WHITE if i == 0 else TEXT)
    return gs


def style_chart(ch, title=None, legend=True, legend_pos=XL_LEGEND_POSITION.BOTTOM, size=13):
    ch.font.name = FONT
    ch.font.size = Pt(size)
    if title:
        ch.has_title = True
        ch.chart_title.text_frame.text = title
        for p in ch.chart_title.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size + 3)
                r.font.bold = True
                r.font.name = FONT
    else:
        ch.has_title = False
    ch.has_legend = legend
    if legend:
        ch.legend.position = legend_pos
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size - 1)
        ch.legend.font.name = FONT


def axis_title(ax, text, size=13):
    ax.has_title = True
    ax.axis_title.text_frame.text = text
    for p in ax.axis_title.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = False
            r.font.name = FONT


def grid(ax, on=True):
    ax.has_major_gridlines = on
    if on:
        ax.major_gridlines.format.line.color.rgb = rgb("D9D9D9")


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", required=True)
    ap.add_argument("--data", required=True)
    ap.add_argument("--assets", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    vres = json.load(open(os.path.join(a.results, "validation", "validation_results.json")))
    dat = vres["data"]
    npass = sum(r["status"] == "PASS" for r in vres["results"])
    ntot = len(vres["results"])
    t3 = {r["lens"]: r for r in dat["T3"]}
    t5 = {r["lens"]: r for r in dat["T5"]}
    o1t = next(r for r in dat["T10"] if r["lens"] == "objective1" and r["E_keV"] == 25)
    wdir = os.path.join(a.data, "05_waist_plane_components")
    waist = {(r["lens"], round(r["energy_keV"], 2)): r for r in rcsv(os.path.join(wdir, "table_waist_and_spot.csv"))}
    comp = {(r["lens"], round(r["energy_keV"], 2)): r for r in rcsv(os.path.join(wdir, "table_waist_plane_components.csv"))
            if r["spot_radius_um"] == 1.0}
    scan = rcsv(os.path.join(a.data, "03_sb_k_edge", "table_sb_edge_scan.csv"))
    sc = {(r["lens"], r["material"], round(r["energy_keV"], 2)): r for r in scan}
    div = rcsv(os.path.join(a.data, "04_focal_spot_divergence", "table_focal_spot_vs_divergence.csv"))
    mu_rows = rcsv(os.path.join(a.data, "00_attenuation", "mu_SU8_15-35keV.csv"))
    mu0_rows = rcsv(os.path.join(a.data, "00_attenuation", "mu_SU8_noSb_15-35keV.csv"))
    mu = {round(r["energy_keV"], 3): r for r in mu_rows}

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(LIGHT)

    # ---------------- header
    rect(slide, 0, 0.912, 1, 0.088, NAVY, rounded=False, name="Шапка")
    textbox(slide, 0.02, 0.995, 0.96, 0.048, ["Моделирование прохождения жёсткого рентгеновского излучения",
                                               "через преломляющие линзы SU-8 микроскопа МИР 15–35 кэВ в Geant4"],
            size=44, color=WHITE, bold=True, align=PP_ALIGN.CENTER, name="Название", spacing=1.0)
    textbox(slide, 0.02, 0.946, 0.96, 0.018, ["А. А. Неустроева¹, Е. Ф. Резникова¹, В. П. Назьмов¹ ² ³"], size=28, color=WHITE,
            bold=True, align=PP_ALIGN.CENTER, name="Авторы")
    textbox(slide, 0.02, 0.928, 0.96, 0.014, ["¹ ИЯФ им. Г. И. Будкера СО РАН   ² ЦКП «СКИФ», ИК СО РАН   ³ ИХТТМ СО РАН,  Новосибирск"],
            size=20, color="DFE8F3", align=PP_ALIGN.CENTER, name="Организации")

    c1, cw1, c2, cw2 = 0.015, 0.30, 0.330, 0.655

    # ---------------- task & model
    section(slide, c1, 0.640, cw1, 0.262, "Задача и модель")
    textbox(slide, c1 + 0.008, 0.876, cw1 - 0.016, 0.07,
            ["Объективы и конденсоры МИР — длинные составные преломляющие линзы из сотен линейных параболических "
             "элементов SU-8, скрещённых по x и y. Цель — модель в Geant4, где преломление на каждой поверхности "
             "считается точно, а в том же материале работают фотоэффект, комптоновское и рэлеевское рассеяние."],
            size=17, name="Текст: задача")
    slide.shapes.add_picture(os.path.join(a.assets, "lens_elements_3d.png"), *box_emu(c1 + 0.01, 0.735, cw1 - 0.02, 0.068))
    textbox(slide, c1 + 0.008, 0.730, cw1 - 0.016, 0.085,
            ["• элементы X, Y, X, Y … (схема, масштаб условный)",
             "• профиль z = ±(d/2 + x²/2R), апертура 2√(R(p − d))",
             "• своё тело G4VSolid: точные пересечения луч–парабола",
             "• процесс преломления γ: векторный закон Снеллиуса в устойчивой форме (δ ~ 10⁻⁷ без потери точности)",
             "• SU-8: 1,24 г/см³, C/O/H/Sb/F/S (Sb 1,33 %), δ(25 кэВ) = 4,32·10⁻⁷; физика Geant4 Livermore + флуоресценция"],
            size=15.5, name="Текст: модель")

    # ---------------- beam picture
    section(slide, c2, 0.640, cw2, 0.262, "Пучок через массив: моделирование Geant4")
    slide.shapes.add_picture(os.path.join(a.assets, "beam_objective1.png"), *box_emu(c2 + 0.008, 0.648, cw2 - 0.016, 0.228))

    # ---------------- stage 1
    section(slide, c1, 0.332, cw1, 0.300, "Этап 1: расчёт = моделирование")
    textbox(slide, c1 + 0.008, 0.606, cw1 - 0.016, 0.012, [[("Фокусное расстояние одного элемента, 25 кэВ", {"bold": True})]],
            size=16)
    rows = [["Один элемент", "R/(2δ), м", "Geant4, м", "отклон."]]
    for l in ("objective1", "objective2", "condenser1", "condenser2"):
        r = t3[l]
        rows.append([NAME[l], rc(r["f_thin_mm"] / 1000, 4), rc(r["efl_A_mm"] / 1000, 4), f"{r['rel_A_thin']:.0e}"])
    table(slide, c1 + 0.008, 0.512, cw1 - 0.016, 0.080, rows, size=15, col_w=[0.34, 0.22, 0.22, 0.2], name="Таблица: один элемент")
    textbox(slide, c1 + 0.008, 0.500, cw1 - 0.016, 0.022,
            [[("Массивы, 25 кэВ: трассировка Geant4 и матричная модель", {"bold": True})], "(фокус — от выходного торца)"], size=16)
    rows = [["Массив", "EFL G4, мм", "EFL матр.", "фокус, мм", "расх."]]
    for l in ("objective1", "objective2_N180", "objective2", "condenser1", "condenser2"):
        r = t5[l]
        rows.append([NAME[l], rc(r["efl_A_x"], 3), rc(r["efl_C_x"], 3), rc(r["bfl_A_x"], 2), f"{abs(r['defl_AC_rel_x']):.0e}"])
    table(slide, c1 + 0.008, 0.383, cw1 - 0.016, 0.090, rows, size=14, col_w=[0.3, 0.19, 0.19, 0.18, 0.14], name="Таблица: массивы")
    textbox(slide, c1 + 0.008, 0.378, cw1 - 0.016, 0.042,
            [f"• {npass}/{ntot} автоматических проверок PASS; закон Снеллиуса: ошибка ≤ 3·10⁻¹⁵ рад",
             f"• табл. 2 диссертации = тонкая линза R/(2Nδ) (±1,2 %); реальный фокус объектива 1: "
             f"{rc(o1t['bfl_end_cm'], 2)} см от торца вместо {rc(o1t['table_cm'], 1)} см"], size=15, name="Текст: этап 1")

    # ---------------- waist plane
    section(slide, c2, 0.332, cw2, 0.300, "Плоскость перетяжки: куда попадают фотоны, упавшие на массив")
    lenses = ["objective1", "objective2_N180", "condenser2"]
    energies = [25.0, 30.4, 30.6, 35.0]
    cd = CategoryChartData()
    for l in lenses:
        grp = cd.add_category(NAME[l])
        for E in energies:
            grp.add_sub_category(rc(E, 1).replace(",0", "") + " кэВ")
    for key, lab in (("pct_in_focal_spot", "в фокусное пятно r ≤ 1 мкм"), ("pct_noise_at_plane", "шум в плоскости перетяжки"),
                     ("pct_absorbed", "поглощены (фотоэффект)"), ("pct_scattered_out_of_beam", "рассеяны из пучка")):
        cd.add_series(lab, [round(comp[(l, E)][key], 1) for l in lenses for E in energies])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, *box_emu(c2 + 0.008, 0.395, 0.455, 0.212), cd)
    gf.name = "Диаграмма: судьба фотонов в плоскости перетяжки"
    ch = gf.chart
    style_chart(ch, "Судьба фотонов (% от упавших на массив), энергия в кэВ")
    plot = ch.plots[0]
    plot.gap_width = 35
    plot.overlap = 100
    for s, c in zip(plot.series, (NAVY, ORANGE, ACC, GREY)):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(c)
        if c != ORANGE:
            s.data_labels.show_value = True
            s.data_labels.number_format = "0"
            s.data_labels.number_format_is_linked = False
            s.data_labels.position = XL_LABEL_POSITION.CENTER
            s.data_labels.font.size = Pt(12)
            s.data_labels.font.bold = True
            s.data_labels.font.color.rgb = rgb(WHITE)
    va = ch.value_axis
    va.maximum_scale = 100
    va.minimum_scale = 0
    grid(va)
    axis_title(va, "% фотонов")
    ch.category_axis.tick_labels.font.size = Pt(11)

    cd = CategoryChartData()
    rows = sorted([r for k, r in comp.items() if k[0] == "objective1"], key=lambda r: r["energy_keV"])
    cd.categories = [rc(r["energy_keV"], 1).replace(",0", "") for r in rows]
    cd.add_series("рэлеевское", [round(100 * (r["n_spot_rayleigh"] + r["n_halo_rayleigh"]) / r["n_incident"], 3) for r in rows])
    cd.add_series("комптоновское", [round(100 * (r["n_spot_compton"] + r["n_halo_compton"]) / r["n_incident"], 3) for r in rows])
    cd.add_series("флуоресценция Sb", [round(100 * r["n_fluorescence_at_plane"] / r["n_incident"], 3) for r in rows])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, *box_emu(c2 + 0.470, 0.470, 0.178, 0.137), cd)
    gf.name = "Диаграмма: состав шума"
    ch = gf.chart
    style_chart(ch, "Состав шума, объектив 1 (% от упавших)", size=12)
    ch.plots[0].gap_width = 50
    for s, c in zip(ch.plots[0].series, (GREEN, ORANGE, VIOLET)):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(c)
    ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 0, 0.3
    ch.value_axis.major_unit = 0.05
    ch.value_axis.tick_labels.number_format = "0.00"
    ch.value_axis.tick_labels.number_format_is_linked = False
    grid(ch.value_axis)
    axis_title(ch.category_axis, "энергия, кэВ", 12)

    w1 = waist[("objective1", 25.0)]
    rows = [["25 кэВ", "перетяжка, мм", "zx − zy, мм", "пятно, нм"]]
    for l in ("objective1", "objective2_N180", "condenser2"):
        w = waist[(l, 25.0)]
        rows.append([NAME[l].replace("Объектив", "Об.").replace("Конденсор", "Конд."), rc(w["waist_from_lens_end_mm"], 1),
                     rc(w["astigmatism_zx_minus_zy_mm"], 2), f"{1000 * w['spot_estimate_thesis_5_2_5_um']:.0f}"])
    textbox(slide, c2 + 0.470, 0.462, 0.178, 0.022, ["Перетяжка — от выходного торца; пятно — оценка по формуле (5.2.5)"],
            size=12, color=MUTED)
    table(slide, c2 + 0.470, 0.352, 0.178, 0.078, rows, size=12, col_w=[0.24, 0.29, 0.25, 0.22], name="Таблица: перетяжка и пятно")
    d1 = next(r for r in div if r["divergence_sigma_urad"] == 1.0)
    textbox(slide, c2 + 0.010, 0.392, 0.452, 0.055,
            ["В фокусное пятно (r ≤ 1 мкм) попадают только фотоны без взаимодействий; рассеянные (0,1–0,8 % упавших) "
             "приходят в сотнях мкм – мм от оси. Геометрическое пятно идеальной линзы < 0,05 мкм (астигматизм X/Y); "
             f"реальное задают дифракция с поглощением (~{1000 * w1['spot_estimate_thesis_5_2_5_um']:.0f} нм по (5.2.5)) "
             f"и источник: при расходимости 1 мкрад σ = {1000 * d1['sigma_x_um']:.0f} нм."], size=15, name="Текст: перетяжка")

    # ---------------- Sb
    section(slide, c1, 0.135, 0.97, 0.190, "Роль сурьмы: K-край 30,49 кэВ", color=SBRED)
    xd = XyChartData()
    def thin(rows_):
        out = []
        for r in rows_:
            e = r["energy_keV"]
            if (abs(e * 2 - round(e * 2)) < 1e-6 and not 30.4 <= e <= 30.6) or e in (30.4, 30.48, 30.49, 30.5, 30.52, 30.6):
                out.append(r)
        return out
    pick, pick0 = thin(mu_rows), thin(mu0_rows)
    for col, lab, rows_ in (("mu_total_per_cm", "SU-8: полное", pick), ("mu_phot_per_cm", "SU-8: фотоэффект", pick),
                            ("mu_phot_Sb_per_cm", "из них сурьма", pick), ("mu_compt_per_cm", "Комптон", pick),
                            ("mu_rayl_per_cm", "Рэлей", pick), ("mu_total_per_cm", "SU-8 без Sb: полное", pick0)):
        s = xd.add_series(lab)
        for r in rows_:
            s.add_data_point(r["energy_keV"], round(r[col], 5))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, *box_emu(c1 + 0.012, 0.142, 0.325, 0.158), xd)
    gf.name = "Диаграмма: μ(E) SU-8"
    ch = gf.chart
    style_chart(ch, "Ослабление в SU-8 (Geant4 Livermore), 1/см", legend_pos=XL_LEGEND_POSITION.BOTTOM, size=12)
    for s, (c, wdt, dash) in zip(ch.plots[0].series, ((NAVY, 3, None), (ACC, 2.25, None), (ACC, 2, "dot"), (ORANGE, 1.75, None),
                                                    (GREEN, 1.75, None), (NAVY, 2.25, "dash"))):
        s.smooth = False
        s.marker.style = XL_MARKER_STYLE.NONE
        s.format.line.color.rgb = rgb(c)
        s.format.line.width = Pt(wdt)
        if dash:
            s.format.line.dash_style = MSO_LINE_DASH_STYLE.SQUARE_DOT if dash == "dot" else MSO_LINE_DASH_STYLE.DASH
    ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 0, 1.6
    ch.value_axis.major_unit = 0.2
    ch.category_axis.minimum_scale, ch.category_axis.maximum_scale = 15, 35
    ch.category_axis.major_unit = 5
    grid(ch.value_axis)
    axis_title(ch.category_axis, "энергия, кэВ", 12)

    xd = XyChartData()
    for lens in ("objective1", "condenser2"):
        for mat in ("SU8", "SU8_noSb"):
            s = xd.add_series(f"{'Об. 1' if lens == 'objective1' else 'Конд. 2'}: {'SU-8' if mat == 'SU8' else 'без Sb'}")
            for r in sorted([r for r in scan if r["lens"] == lens and r["material"] == mat], key=lambda r: r["energy_keV"]):
                s.add_data_point(r["energy_keV"], round(100 * r["frac_unscattered"], 2))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, *box_emu(c1 + 0.345, 0.142, 0.310, 0.158), xd)
    gf.name = "Диаграмма: пропускание через K-край"
    ch = gf.chart
    style_chart(ch, "Прошли без взаимодействий, %", legend_pos=XL_LEGEND_POSITION.BOTTOM, size=12)
    for s, (c, dash, mk) in zip(ch.plots[0].series, ((NAVY, False, XL_MARKER_STYLE.CIRCLE), (NAVY, True, XL_MARKER_STYLE.SQUARE),
                                                   (GREEN, False, XL_MARKER_STYLE.CIRCLE), (GREEN, True, XL_MARKER_STYLE.SQUARE))):
        s.smooth = False
        s.format.line.color.rgb = rgb(c)
        s.format.line.width = Pt(2)
        if dash:
            s.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        s.marker.style = mk
        s.marker.size = 6
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = rgb(c)
        s.marker.format.line.color.rgb = rgb(c)
    ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 0, 60
    ch.value_axis.major_unit = 10
    ch.category_axis.minimum_scale, ch.category_axis.maximum_scale = 25, 35
    ch.category_axis.major_unit = 1
    grid(ch.value_axis)
    axis_title(ch.category_axis, "энергия, кэВ", 12)

    m25, m30a, m30b = mu[25.0], mu[30.48], mu[30.5]
    lo, hi = sc[("objective1", "SU8", 30.48)], sc[("objective1", "SU8", 30.5)]
    ns, s25, s35 = sc[("objective1", "SU8_noSb", 25.0)], sc[("objective1", "SU8", 25.0)], sc[("objective1", "SU8", 35.0)]
    textbox(slide, c1 + 0.665, 0.300, 0.298, 0.155,
            [f"• 1,33 мас.% Sb дают {100 * m25['mu_phot_Sb_per_cm'] / m25['mu_phot_per_cm']:.0f}% фотопоглощения SU-8 уже при 25 кэВ",
             f"• на краю μ: {rc(m30a['mu_total_per_cm'], 3)} → {rc(m30b['mu_total_per_cm'], 3)} 1/см "
             f"(×{rc(m30b['mu_total_per_cm'] / m30a['mu_total_per_cm'], 1)})",
             f"• объектив 1, без взаимодействий: {100 * lo['frac_unscattered']:.0f}% (30,48 кэВ) → {100 * hi['frac_unscattered']:.0f}% "
             f"(30,50 кэВ); при 35 кэВ — {100 * s35['frac_unscattered']:.0f}%",
             f"• без Sb при 25 кэВ: {100 * ns['frac_unscattered']:.0f}% вместо {100 * s25['frac_unscattered']:.0f}%",
             "• выше края — флуоресценция Sb Kα/Kβ (вне фокусного пятна)"], size=18, name="Текст: сурьма")

    # ---------------- conclusions
    section(slide, c1, 0.012, 0.97, 0.116, "Выводы")
    textbox(slide, c1 + 0.010, 0.104, 0.95, 0.088,
            ["1. Модель Geant4 с точным преломлением воспроизводит аналитику: один элемент — 10⁻⁷, массивы из 300–430 элементов — "
             "10⁻¹⁰ (матричная модель).",
             "2. Табл. 2 — тонколинзовые оценки: фокус длинных объективов в 1,8–3 раза ближе к торцу; объектив 2 (N = 215) при 25 кэВ "
             "фокусирует внутри себя.",
             f"3. В фокус приходит только непровзаимодействовавшая часть пучка: {comp[('condenser1', 25.0)]['pct_in_focal_spot']:.0f}–"
             f"{comp[('objective2_N145', 25.0)]['pct_in_focal_spot']:.0f}% при 25 кэВ; остальное поглощается (фотоэффект) или "
             "рассеивается из пучка; шум в плоскости перетяжки < 1 % и весь вне пятна.",
             "4. Над K-краем Sb пропускание падает вдвое: рабочую энергию МИР выгодно держать ≤ 30,4 кэВ, а снижение доли Sb в резисте "
             "увеличит пропускание (без Sb при 25 кэВ — в 1,3 раза больше фотонов в фокусе).",
             "5. Монте-Карло совпадает с законом Бугера–Ламберта–Бера по преломлённым траекториям; формула (5.2.6) с прямыми лучами "
             "занижает пропускание длинных объективов на 18–37 %, для конденсоров верна до 1 %."],
            size=17, name="Текст: выводы", spacing=1.1)
    textbox(slide, 0.30, 0.011, 0.685, 0.008, ["Geant4 11.4.2 · Монте-Карло 10⁵ фотонов на точку · все графики построены из файлов "
                                              "данных проекта"], size=12, color=MUTED, align=PP_ALIGN.RIGHT)
    prs.save(a.out)
    print("pptx written:", a.out)


if __name__ == "__main__":
    main()
